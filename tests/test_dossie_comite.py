"""Testes unitários e de integração para o motor do Dossiê do Comitê de Crédito."""
import io
import openpyxl
import pytest

import database as db
from services.dossie_comite import (
    ConsolidadoAtividades,
    ConsolidadoESG,
    ConsolidadoFinanceiro,
    ConsolidadoGarantias,
    DeliberacaoComite,
    DossieComite,
    gerar_excel_dossie_comite,
    gerar_pdf_dossie_comite,
    montar_dossie_comite,
)


@pytest.fixture
def payload_misto():
    return {
        'identificacao': {
            'nome': 'João da Silva',
            'cpf_cnpj': '123.456.789-00',
            'fazenda': 'Fazenda Primavera',
            'municipio': 'Sorriso',
            'uf': 'MT',
            'area_total_ha': 2500.0,
        },
        'operacao': {
            'valor_solicitado': 2_000_000.0,
            'prazo_meses': 36,
            'carencia_meses': 12,
            'juros_aa': 11.5,
            'sistema_amortizacao': 'PRICE',
            'finalidade': 'Custeio Agrícola e Reforma de Pastagem B2B',
        },
        'pecuaria': {
            'rebanho_total_cabecas': 1200,
            'rebanho_uas': 900.0,
            'taxa_desfrute_pct': 24.5,
            'receita_pecuaria_anual': 1_800_000.0,
        },
        'agricola': {
            'area_plantada_ha': 1000.0,
            'culturas_principais': ['Soja 1ª Safra', 'Milho Safrinha'],
            'produtividade_media_sc_ha': 62.0,
            'breakeven_medio_sc_ha': 42.5,
            'receita_agricola_anual': 4_500_000.0,
            'economico': {
                'cot_total': 3_200_000.0,
                'receita_total': 4_500_000.0,
            }
        },
        'esg': {
            'numero_car': 'MT-5107909-67D890E1F2A3B4C5D6E7F8A9B0C1D2E3',
            'status_car': 'ATIVO',
            'bioma': 'CERRADO',
            'area_total_ha': 2500.0,
            'area_reserva_legal_ha': 500.0,
            'area_app_ha': 150.0,
            'status_ibama': 'REGULAR',
            'status_icmbio': 'REGULAR',
            'sobreposicao_indigena_quilombola': False,
            'parecer': 'APROVADO',
        },
        'garantias': {
            'valor_mercado_total': 12_000_000.0,
            'valor_liquidacao_total': 9_600_000.0,
            'itens': [
                {
                    'tipo': 'IMOVEL_RURAL',
                    'descricao': 'Fazenda Primavera - Matrícula 9876 CRI Sorriso',
                    'gravame': 'ALIENACAO_FIDUCIARIA',
                    'valor_mercado_bruto': 10_000_000.0,
                    'desagio_aplicado_pct': 20.0,
                    'valor_liquidacao_forcada': 8_000_000.0,
                },
                {
                    'tipo': 'PENHOR_SAFRA',
                    'descricao': 'CPR Física - Penhor de Soja Safra 2026/27',
                    'gravame': 'PENHOR_PRIMEIRO_GRAU',
                    'valor_mercado_bruto': 2_000_000.0,
                    'desagio_aplicado_pct': 20.0,
                    'valor_liquidacao_forcada': 1_600_000.0,
                }
            ]
        }
    }


def test_montar_dossie_misto(payload_misto):
    dossie = montar_dossie_comite(payload_misto)
    assert isinstance(dossie, DossieComite)

    # Proponente
    assert dossie.proponente.nome == 'João da Silva'
    assert dossie.proponente.cpf_cnpj == '123.456.789-00'
    assert dossie.proponente.area_total_ha == 2500.0

    # Operação
    assert dossie.operacao.valor_solicitado == 2_000_000.0
    assert dossie.operacao.prazo_meses == 36
    assert dossie.operacao.codigo_operacao.startswith('ORK-COM-')

    # Atividades
    assert dossie.atividades.tem_pecuaria is True
    assert dossie.atividades.rebanho_total_cabecas == 1200
    assert dossie.atividades.tem_agricola is True
    assert dossie.atividades.area_plantada_ha == 1000.0

    # ESG
    assert dossie.esg.parecer_compliance == 'APROVADO'
    assert dossie.esg.status_car == 'ATIVO'
    assert dossie.esg.reserva_legal_pct == 20.0

    # Garantias
    assert dossie.garantias.total_bens == 2
    assert dossie.garantias.valor_mercado_total == 12_000_000.0
    assert dossie.garantias.valor_liquidacao_forcada_total == 9_600_000.0
    assert dossie.garantias.ltv_liquidacao_forcada_pct < 30.0
    assert dossie.garantias.faixa_risco_ltv == 'CONFORTAVEL'

    # Financeiro
    assert dossie.financeiro.receita_bruta_anual == 6_300_000.0  # 1.8M + 4.5M
    assert dossie.financeiro.ebitda_anual > 0
    assert dossie.financeiro.dscr_medio > 1.30
    assert len(dossie.financeiro.projecao_plurianual) == 3

    # Deliberação
    assert dossie.deliberacao.recomendacao == 'FAVORAVEL'
    assert dossie.deliberacao.limite_sugerido == 2_000_000.0
    assert len(dossie.deliberacao.covenants_obrigatorios) >= 4
    assert len(dossie.deliberacao.hash_sha256) == 64


def test_montar_dossie_pecuaria_pura():
    payload = {
        'identificacao': {'fazenda': 'Estância Nelore', 'proprietario': 'Carlos Nelore'},
        'operacao': {'valor_solicitado': 500_000.0, 'prazo_meses': 24},
        'saldo': {'total': 800},
        'fluxo_gep': {'receita_total': 1_200_000.0, 'custo_total': 700_000.0},
    }
    dossie = montar_dossie_comite(payload)
    assert dossie.atividades.tem_pecuaria is True
    assert dossie.atividades.rebanho_total_cabecas == 800
    assert dossie.atividades.tem_agricola is False
    assert dossie.financeiro.receita_bruta_anual == 1_200_000.0
    assert dossie.financeiro.ebitda_anual == 500_000.0
    assert dossie.deliberacao.recomendacao in ('FAVORAVEL', 'FAVORAVEL_COM_RESSALVAS')


def test_montar_dossie_agricola_pura():
    payload = {
        'identificacao': {'fazenda': 'Fazenda Soja Ouro', 'proprietario': 'Marcos Grãos'},
        'operacao': {'valor_solicitado': 1_500_000.0, 'prazo_meses': 24},
        'economico': {
            'receita_total': 3_800_000.0,
            'cot_total': 2_400_000.0,
            'culturas': [{'cultura': 'SOJA', 'area_ha': 600, 'produtividade_sc_ha': 65, 'breakeven_sc_ha': 41}]
        }
    }
    dossie = montar_dossie_comite(payload)
    assert dossie.atividades.tem_agricola is True
    assert dossie.atividades.tem_pecuaria is False
    assert dossie.atividades.area_plantada_ha == 600.0
    assert dossie.financeiro.receita_bruta_anual == 3_800_000.0


def test_deliberacao_bloqueio_esg(payload_misto):
    payload_misto['esg']['parecer'] = 'BLOQUEIO'
    dossie = montar_dossie_comite(payload_misto)
    assert dossie.deliberacao.recomendacao == 'DESFAVORAVEL'
    assert dossie.deliberacao.limite_sugerido == 0.0
    assert 'socioambiental' in dossie.deliberacao.parecer_risco_tecnico.lower()


def test_deliberacao_dscr_insuficiente(payload_misto):
    # Pedir 50M de crédito para uma receita de 6M -> DSCR colapsa
    payload_misto['operacao']['valor_solicitado'] = 50_000_000.0
    dossie = montar_dossie_comite(payload_misto)
    assert dossie.financeiro.dscr_minimo < 1.0
    assert dossie.deliberacao.recomendacao == 'DESFAVORAVEL'
    assert dossie.deliberacao.limite_sugerido == 0.0


def test_gerar_pdf_dossie_comite(payload_misto):
    dossie = montar_dossie_comite(payload_misto)
    pdf_bytes = gerar_pdf_dossie_comite(dossie)
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 2000
    assert pdf_bytes.startswith(b'%PDF-')


def test_gerar_pdf_dossie_comite_com_dict(payload_misto):
    dossie_dict = montar_dossie_comite(payload_misto).to_dict()
    pdf_bytes = gerar_pdf_dossie_comite(dossie_dict, branding={'nome_consultoria': 'AgroCapital B2B'})
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 2000
    assert pdf_bytes.startswith(b'%PDF-')


def test_gerar_excel_dossie_comite(payload_misto):
    dossie = montar_dossie_comite(payload_misto)
    xlsx_bytes = gerar_excel_dossie_comite(dossie)
    assert isinstance(xlsx_bytes, bytes)
    assert len(xlsx_bytes) > 2000

    # Carrega e valida via openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    abas_esperadas = [
        "1. Resumo do Comitê",
        "2. Fluxo Caixa Consolidado",
        "3. Matriz de Garantias",
        "4. Compliance ESG",
        "5. Atividades Agropecuárias"
    ]
    assert wb.sheetnames == abas_esperadas

    ws1 = wb["1. Resumo do Comitê"]
    assert "ORKAVYN AGRO INTELLIGENCE" in str(ws1['A1'].value)
    assert "DADOS DO PROPONENTE" in str(ws1['A7'].value)

    ws2 = wb["2. Fluxo Caixa Consolidado"]
    assert ws2['A3'].value == 'Ano'
    assert ws2['G3'].value == 'DSCR (x)'
    assert ws2['A4'].value == 'Ano 1'

    ws3 = wb["3. Matriz de Garantias"]
    assert ws3['A3'].value == 'Garantia / Descrição'

    ws4 = wb["4. Compliance ESG"]
    assert ws4['A1'].value.startswith('CONFORMIDADE SOCIOAMBIENTAL')

    ws5 = wb["5. Atividades Agropecuárias"]
    assert ws5['A1'].value.startswith('BASE PRODUTIVA')


def test_api_rotas_comite(payload_misto):
    from app import app
    app.config['TESTING'] = True

    email = 'analista_comite@orkavyn.test'
    usuario = db.buscar_usuario_email(email)
    if not usuario:
        db.criar_usuario(email, 'Hash123', 'Analista Comite', 'USER')
        usuario = db.buscar_usuario_email(email)

    cliente = app.test_client()
    with cliente.session_transaction() as sessao:
        sessao['_user_id'] = str(usuario['id'])

    # 1. Rota de Consolidação
    resp_cons = cliente.post('/api/comite/dossie/consolidar', json=payload_misto)
    assert resp_cons.status_code == 200
    dados_dossie = resp_cons.get_json()
    assert 'deliberacao' in dados_dossie
    assert dados_dossie['proponente']['nome'] == 'João da Silva'
    assert dados_dossie['deliberacao']['rating_final'] in ('AAA', 'AA', 'A', 'BBB')

    # 2. Rota de PDF
    resp_pdf = cliente.post('/api/comite/dossie/pdf', json=dados_dossie)
    assert resp_pdf.status_code == 200
    assert resp_pdf.content_type == 'application/pdf'
    assert resp_pdf.data.startswith(b'%PDF-')

    # 3. Rota de Excel
    resp_xls = cliente.post('/api/comite/dossie/excel', json=dados_dossie)
    assert resp_xls.status_code == 200
    assert 'spreadsheet' in resp_xls.content_type
    wb = openpyxl.load_workbook(io.BytesIO(resp_xls.data))
    assert len(wb.sheetnames) == 5
