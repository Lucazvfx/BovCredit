"""Gerador de planilha Excel (.xlsx) profissional para o Dossiê do Comitê de Crédito."""
from __future__ import annotations

import io
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from services.dossie_comite.models import DossieComite

# Paleta Institucional Orkavyn
_HEX_VERDE_ESCURO = '1E3A2F'
_HEX_VERDE_MEDIO = '2D5A46'
_HEX_BEGE = 'F3EFE6'
_HEX_CINZA_CLARO = 'F8F9FA'
_HEX_BORDA = 'D8C9B3'
_HEX_BRANCO = 'FFFFFF'

_FONT_TITULO = Font(name='Calibri', size=14, bold=True, color=_HEX_VERDE_ESCURO)
_FONT_SECAO = Font(name='Calibri', size=11, bold=True, color=_HEX_BRANCO)
_FONT_SUBSECAO = Font(name='Calibri', size=11, bold=True, color=_HEX_VERDE_ESCURO)
_FONT_HEADER = Font(name='Calibri', size=10, bold=True, color=_HEX_BRANCO)
_FONT_DADO = Font(name='Calibri', size=10)
_FONT_DADO_BOLD = Font(name='Calibri', size=10, bold=True)
_FONT_NOTA = Font(name='Calibri', size=9, italic=True, color='666666')

_FILL_HEADER = PatternFill(start_color=_HEX_VERDE_ESCURO, end_color=_HEX_VERDE_ESCURO, fill_type='solid')
_FILL_SUBHEADER = PatternFill(start_color=_HEX_BEGE, end_color=_HEX_BEGE, fill_type='solid')
_FILL_ZEBRA = PatternFill(start_color=_HEX_CINZA_CLARO, end_color=_HEX_CINZA_CLARO, fill_type='solid')

_BORDER_THIN = Border(
    left=Side(style='thin', color=_HEX_BORDA),
    right=Side(style='thin', color=_HEX_BORDA),
    top=Side(style='thin', color=_HEX_BORDA),
    bottom=Side(style='thin', color=_HEX_BORDA),
)


def _ajustar_larguras(ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = cell.value
            if val is not None:
                tam = len(str(val))
                if tam > max_len and tam < 60:
                    max_len = tam
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)


def gerar_excel_dossie_comite(dossie: DossieComite | dict[str, Any]) -> bytes:
    """Gera uma pasta de trabalho Excel (.xlsx) com modelagem completa do Comitê."""
    if isinstance(dossie, DossieComite):
        d = dossie.to_dict()
    else:
        d = dossie

    proponente = d.get('proponente') or {}
    operacao = d.get('operacao') or {}
    atividades = d.get('atividades') or {}
    esg = d.get('esg') or {}
    garantias = d.get('garantias') or {}
    financeiro = d.get('financeiro') or {}
    deliberacao = d.get('deliberacao') or {}

    wb = openpyxl.Workbook()

    # ─────────────────────────────────────────────────────────────
    # ABA 1: RESUMO COMITÊ
    # ─────────────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "1. Resumo do Comitê"
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells('A1:F1')
    ws1['A1'] = "ORKAVYN AGRO INTELLIGENCE — DOSSIÊ DE CRÉDITO B2B"
    ws1['A1'].font = _FONT_TITULO
    ws1['A1'].alignment = Alignment(horizontal='left', vertical='center')

    ws1.merge_cells('A2:F2')
    ws1['A2'] = f"Operação: {deliberacao.get('codigo_dossie')} · Emissão: {deliberacao.get('data_emissao')}"
    ws1['A2'].font = _FONT_NOTA

    # Fact Sheet Card em Destaque
    r = 4
    cards = [
        ('RATING FINAL', deliberacao.get('rating_final')),
        ('RECOMENDAÇÃO', str(deliberacao.get('recomendacao')).replace('_', ' ')),
        ('DSCR MÉDIO', f"{financeiro.get('dscr_medio'):.2f}x"),
        ('LTV LIQUIDAÇÃO', f"{garantias.get('ltv_liquidacao_forcada_pct'):.1f}%"),
        ('COMPLIANCE ESG', esg.get('parecer_compliance')),
    ]
    for idx, (label, val) in enumerate(cards, start=1):
        cell_lbl = ws1.cell(row=r, column=idx, value=label)
        cell_lbl.font = _FONT_HEADER
        cell_lbl.fill = _FILL_HEADER
        cell_lbl.alignment = Alignment(horizontal='center', vertical='center')
        cell_lbl.border = _BORDER_THIN

        cell_val = ws1.cell(row=r + 1, column=idx, value=val)
        cell_val.font = Font(name='Calibri', size=12, bold=True, color=_HEX_VERDE_ESCURO)
        cell_val.fill = _FILL_SUBHEADER
        cell_val.alignment = Alignment(horizontal='center', vertical='center')
        cell_val.border = _BORDER_THIN

    # Proponente e Operação
    r = 7
    ws1.merge_cells(f'A{r}:F{r}')
    ws1[f'A{r}'] = "DADOS DO PROPONENTE & OPERAÇÃO DE CRÉDITO"
    ws1[f'A{r}'].font = _FONT_SECAO
    ws1[f'A{r}'].fill = _FILL_HEADER

    linhas_cad = [
        ("Proponente:", proponente.get('nome'), "Valor Solicitado:", operacao.get('valor_solicitado')),
        ("CPF / CNPJ:", proponente.get('cpf_cnpj'), "Prazo / Carência:", f"{operacao.get('prazo_meses')}m / {operacao.get('carencia_meses')}m"),
        ("Fazenda:", proponente.get('fazenda'), "Taxa de Juros:", f"{operacao.get('juros_aa')}% a.a."),
        ("Município / UF:", f"{proponente.get('municipio')}/{proponente.get('uf')}", "Amortização:", f"{operacao.get('sistema_amortizacao')} / {operacao.get('periodicidade')}"),
        ("Área Total (ha):", proponente.get('area_total_ha'), "Finalidade:", operacao.get('finalidade')),
    ]
    for c1, v1, c2, v2 in linhas_cad:
        r += 1
        ws1.cell(row=r, column=1, value=c1).font = _FONT_DADO_BOLD
        c_v1 = ws1.cell(row=r, column=2, value=v1)
        c_v1.font = _FONT_DADO
        ws1.cell(row=r, column=4, value=c2).font = _FONT_DADO_BOLD
        c_v2 = ws1.cell(row=r, column=5, value=v2)
        c_v2.font = _FONT_DADO
        if c2 == "Valor Solicitado:":
            c_v2.number_format = '"R$" #,##0.00'

    # Covenants
    r += 2
    ws1.merge_cells(f'A{r}:F{r}')
    ws1[f'A{r}'] = "COVENANTS OBRIGATÓRIOS & CONDICIONANTES DO FINANCIAMENTO"
    ws1[f'A{r}'].font = _FONT_SECAO
    ws1[f'A{r}'].fill = _FILL_HEADER

    for cov in deliberacao.get('covenants_obrigatorios') or []:
        r += 1
        ws1.merge_cells(f'A{r}:F{r}')
        ws1[f'A{r}'] = f"• {cov}"
        ws1[f'A{r}'].font = _FONT_DADO

    # Hash de Integridade
    r += 2
    ws1.merge_cells(f'A{r}:F{r}')
    ws1[f'A{r}'] = f"Hash Criptográfico SHA-256: {deliberacao.get('hash_sha256')}"
    ws1[f'A{r}'].font = Font(name='Courier New', size=9, bold=True, color='444444')

    _ajustar_larguras(ws1)

    # ─────────────────────────────────────────────────────────────
    # ABA 2: FLUXO DE CAIXA CONSOLIDADO
    # ─────────────────────────────────────────────────────────────
    ws2 = wb.create_sheet(title="2. Fluxo Caixa Consolidado")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells('A1:G1')
    ws2['A1'] = "PROJEÇÃO DE FLUXO DE CAIXA LIVRE E CAPACIDADE DE PAGAMENTO (DSCR)"
    ws2['A1'].font = _FONT_TITULO

    headers_fluxo = ['Ano', 'Receita Bruta (R$)', 'Custos Operacionais (R$)', 'EBITDA (R$)', 'Serviço Dívida (R$)', 'Fluxo Caixa Livre (R$)', 'DSCR (x)']
    for c_idx, h in enumerate(headers_fluxo, start=1):
        cell = ws2.cell(row=3, column=c_idx, value=h)
        cell.font = _FONT_HEADER
        cell.fill = _FILL_HEADER
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = _BORDER_THIN

    r = 4
    for l in (financeiro.get('projecao_plurianual') or []):
        ws2.cell(row=r, column=1, value=f"Ano {l.get('ano')}").alignment = Alignment(horizontal='center')
        ws2.cell(row=r, column=2, value=l.get('receita_bruta')).number_format = '"R$" #,##0.00'
        ws2.cell(row=r, column=3, value=l.get('custos_despesas')).number_format = '"R$" #,##0.00'
        ws2.cell(row=r, column=4, value=l.get('ebitda')).number_format = '"R$" #,##0.00'
        ws2.cell(row=r, column=5, value=l.get('servico_divida')).number_format = '"R$" #,##0.00'
        ws2.cell(row=r, column=6, value=l.get('fluxo_caixa_livre')).number_format = '"R$" #,##0.00'
        ws2.cell(row=r, column=7, value=l.get('dscr')).number_format = '0.00"x"'

        for c_idx in range(1, 8):
            cell = ws2.cell(row=r, column=c_idx)
            cell.font = _FONT_DADO
            cell.border = _BORDER_THIN
            if r % 2 == 0:
                cell.fill = _FILL_ZEBRA
        r += 1

    r += 1
    ws2.cell(row=r, column=1, value="MÉTRICAS:").font = _FONT_DADO_BOLD
    ws2.cell(row=r, column=2, value=f"EBITDA Base: R$ {financeiro.get('ebitda_anual'):,.2f}").font = _FONT_DADO_BOLD
    ws2.cell(row=r, column=4, value=f"DSCR Mínimo: {financeiro.get('dscr_minimo'):.2f}x (Ano {financeiro.get('ano_critico')})").font = _FONT_DADO_BOLD
    ws2.cell(row=r, column=6, value=f"DSCR Médio: {financeiro.get('dscr_medio'):.2f}x").font = _FONT_DADO_BOLD

    _ajustar_larguras(ws2)

    # ─────────────────────────────────────────────────────────────
    # ABA 3: MATRIZ DE GARANTIAS & LTV
    # ─────────────────────────────────────────────────────────────
    ws3 = wb.create_sheet(title="3. Matriz de Garantias")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells('A1:F1')
    ws3['A1'] = "MATRIZ DE GARANTIAS, GRAVAMES E COBERTURA DE LTV B2B"
    ws3['A1'].font = _FONT_TITULO

    headers_gar = ['Garantia / Descrição', 'Tipo Gravame', 'Valor Mercado (R$)', 'Deságio (%)', 'Liq. Forçada (R$)', 'LTV Ponderado']
    for c_idx, h in enumerate(headers_gar, start=1):
        cell = ws3.cell(row=3, column=c_idx, value=h)
        cell.font = _FONT_HEADER
        cell.fill = _FILL_HEADER
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = _BORDER_THIN

    r = 4
    for item in (garantias.get('detalhes_itens') or []):
        ws3.cell(row=r, column=1, value=item.get('descricao') or item.get('tipo'))
        ws3.cell(row=r, column=2, value=str(item.get('gravame') or 'ALIENACAO_FIDUCIARIA'))
        ws3.cell(row=r, column=3, value=float(item.get('valor_mercado_bruto') or 0.0)).number_format = '"R$" #,##0.00'
        ws3.cell(row=r, column=4, value=float(item.get('desagio_aplicado_pct') or 0.0) / 100.0).number_format = '0.0%'
        ws3.cell(row=r, column=5, value=float(item.get('valor_liquidacao_forcada') or 0.0)).number_format = '"R$" #,##0.00'
        ws3.cell(row=r, column=6, value=garantias.get('faixa_risco_ltv'))

        for c_idx in range(1, 7):
            cell = ws3.cell(row=r, column=c_idx)
            cell.font = _FONT_DADO
            cell.border = _BORDER_THIN
        r += 1

    # Totais de Garantia
    ws3.cell(row=r, column=1, value="TOTAL CONSOLIDADO").font = _FONT_DADO_BOLD
    ws3.cell(row=r, column=3, value=garantias.get('valor_mercado_total')).number_format = '"R$" #,##0.00'
    ws3.cell(row=r, column=3).font = _FONT_DADO_BOLD
    ws3.cell(row=r, column=5, value=garantias.get('valor_liquidacao_forcada_total')).number_format = '"R$" #,##0.00'
    ws3.cell(row=r, column=5).font = _FONT_DADO_BOLD

    r += 2
    ws3.cell(row=r, column=1, value="LTV Mercado:").font = _FONT_DADO_BOLD
    ws3.cell(row=r, column=2, value=f"{garantias.get('ltv_mercado_pct'):.1f}%").font = _FONT_DADO_BOLD
    ws3.cell(row=r, column=4, value="LTV Liquidação Forçada:").font = _FONT_DADO_BOLD
    ws3.cell(row=r, column=5, value=f"{garantias.get('ltv_liquidacao_forcada_pct'):.1f}%").font = _FONT_DADO_BOLD

    _ajustar_larguras(ws3)

    # ─────────────────────────────────────────────────────────────
    # ABA 4: COMPLIANCE ESG
    # ─────────────────────────────────────────────────────────────
    ws4 = wb.create_sheet(title="4. Compliance ESG")
    ws4.views.sheetView[0].showGridLines = True

    ws4.merge_cells('A1:D1')
    ws4['A1'] = "CONFORMIDADE SOCIOAMBIENTAL & CAR (CMN 5.081/23 E MCR 2-1)"
    ws4['A1'].font = _FONT_TITULO

    itens_esg = [
        ("Inscrição no SICAR (CAR)", esg.get('numero_car')),
        ("Status no SICAR", esg.get('status_car')),
        ("Bioma de Referência", esg.get('bioma')),
        ("Área Total da Propriedade (ha)", esg.get('area_total_ha')),
        ("Reserva Legal Declarada (ha)", esg.get('reserva_legal_ha')),
        ("Percentual Reserva Legal", f"{esg.get('reserva_legal_pct')}% (Mínimo exigido: {esg.get('exigencia_reserva_legal_pct')}%)"),
        ("Área de Preservação Permanente (ha)", esg.get('app_ha')),
        ("Status IBAMA", esg.get('status_ibama')),
        ("Status ICMBio", esg.get('status_icmbio')),
        ("Sobreposição com Terras Indígenas / Quilombolas", "SIM (IMPEDITIVO)" if esg.get('sobreposicao_indigena_quilombola') else "NÃO"),
        ("PARECER FINAL ESG", esg.get('parecer_compliance')),
    ]
    r = 3
    for k, v in itens_esg:
        ws4.cell(row=r, column=1, value=k).font = _FONT_DADO_BOLD
        ws4.cell(row=r, column=1).border = _BORDER_THIN
        ws4.cell(row=r, column=2, value=v).font = _FONT_DADO
        ws4.cell(row=r, column=2).border = _BORDER_THIN
        if k == "PARECER FINAL ESG":
            ws4.cell(row=r, column=2).font = _FONT_DADO_BOLD
            ws4.cell(row=r, column=2).fill = _FILL_SUBHEADER
        r += 1

    r += 1
    ws4.merge_cells(f'A{r}:D{r}')
    ws4[f'A{r}'] = "Observações de Conformidade:"
    ws4[f'A{r}'].font = _FONT_DADO_BOLD
    for obs in esg.get('observacoes') or []:
        r += 1
        ws4.merge_cells(f'A{r}:D{r}')
        ws4[f'A{r}'] = f"• {obs}"
        ws4[f'A{r}'].font = _FONT_DADO

    _ajustar_larguras(ws4)

    # ─────────────────────────────────────────────────────────────
    # ABA 5: ATIVIDADES AGROPECUÁRIAS
    # ─────────────────────────────────────────────────────────────
    ws5 = wb.create_sheet(title="5. Atividades Agropecuárias")
    ws5.views.sheetView[0].showGridLines = True

    ws5.merge_cells('A1:E1')
    ws5['A1'] = "BASE PRODUTIVA E CAPACIDADE ZOOTÉCNICA / AGRÍCOLA"
    ws5['A1'].font = _FONT_TITULO

    r = 3
    ws5.merge_cells(f'A{r}:E{r}')
    ws5[f'A{r}'] = "1. Pecuária de Corte / Leite"
    ws5[f'A{r}'].font = _FONT_SECAO
    ws5[f'A{r}'].fill = _FILL_HEADER

    itens_pec = [
        ("Atividade Pecuária Ativa:", "SIM" if atividades.get('tem_pecuaria') else "NÃO"),
        ("Rebanho Total Declarado:", f"{atividades.get('rebanho_total_cabecas'):,} cabeças"),
        ("Efetivo em Unidades Animais (UA):", f"{atividades.get('rebanho_uas'):,.1f} UA"),
        ("Taxa de Desfrute Estimada:", f"{atividades.get('taxa_desfrute_pct'):.1f}%"),
        ("Receita Pecuária Bruta Anual:", atividades.get('receita_pecuaria_anual')),
    ]
    for k, v in itens_pec:
        r += 1
        ws5.cell(row=r, column=1, value=k).font = _FONT_DADO_BOLD
        cell_v = ws5.cell(row=r, column=2, value=v)
        cell_v.font = _FONT_DADO
        if k == "Receita Pecuária Bruta Anual:":
            cell_v.number_format = '"R$" #,##0.00'

    r += 2
    ws5.merge_cells(f'A{r}:E{r}')
    ws5[f'A{r}'] = "2. Agricultura / Grãos e Perenes"
    ws5[f'A{r}'].font = _FONT_SECAO
    ws5[f'A{r}'].fill = _FILL_HEADER

    cults_str = ", ".join(atividades.get('culturas_principais') or []) or 'N/A'
    itens_agr = [
        ("Atividade Agrícola Ativa:", "SIM" if atividades.get('tem_agricola') else "NÃO"),
        ("Culturas Exploradas:", cults_str),
        ("Área Plantada Total (ha):", f"{atividades.get('area_plantada_ha'):,.1f} ha"),
        ("Produtividade Média:", f"{atividades.get('produtividade_media_sc_ha'):.1f} sc/ha"),
        ("Breakeven Operacional Médio:", f"{atividades.get('breakeven_medio_sc_ha'):.1f} sc/ha"),
        ("Receita Agrícola Bruta Anual:", atividades.get('receita_agricola_anual')),
    ]
    for k, v in itens_agr:
        r += 1
        ws5.cell(row=r, column=1, value=k).font = _FONT_DADO_BOLD
        cell_v = ws5.cell(row=r, column=2, value=v)
        cell_v.font = _FONT_DADO
        if k == "Receita Agrícola Bruta Anual:":
            cell_v.number_format = '"R$" #,##0.00'

    _ajustar_larguras(ws5)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
